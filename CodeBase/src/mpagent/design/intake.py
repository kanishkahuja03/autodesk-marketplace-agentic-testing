"""Stage 0: Intake.

Takes an unstructured input folder and produces:
  - documents/D1_<name>.txt, D2_... : plain text extracted from every document
  - intake_manifest.json            : list of files, hashes, installer facts, warnings

Only the input files are read. Nothing is installed or executed, and the machine is not inspected.
"""

import codecs
import hashlib
import os
import re
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from mpagent.models import Document, FileKind, InputFile, InstallerInfo, IntakeManifest

DOCUMENT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".html", ".htm"}
METADATA_FILENAME = "metadata.json"

# Rough budget for all document text sent to the LLM in one request.
MAX_ESTIMATED_TOKENS = 200_000
# Fewer characters than this per PDF page suggests a scanned (image-only) PDF.
MIN_CHARS_PER_PDF_PAGE = 50


class IntakeError(Exception):
    """The input folder can't be processed (e.g. no installer)."""


# ---------------------------------------------------------------- classification


def classify(path: Path) -> FileKind:
    """Decide what a file is from its name/extension (and, for .zip, its contents)."""
    name = path.name.lower()
    ext = path.suffix.lower()
    if name == METADATA_FILENAME:
        return FileKind.METADATA
    if ext in DOCUMENT_EXTENSIONS:
        return FileKind.DOCUMENT
    if ext in {".msi", ".exe"}:
        return FileKind.INSTALLER
    if ext == ".zip":
        return FileKind.INSTALLER if _zip_contains_bundle(path) else FileKind.UNSUPPORTED
    return FileKind.UNSUPPORTED


def _zip_contains_bundle(path: Path) -> bool:
    """An Autodesk plug-in zip contains a *.bundle folder with PackageContents.xml."""
    try:
        with zipfile.ZipFile(path) as zf:
            return any(n.lower().endswith("packagecontents.xml") for n in zf.namelist())
    except zipfile.BadZipFile:
        return False


# ---------------------------------------------------------------- hashing


def sha256_of(path: Path) -> str:
    """Fingerprint of the file; changes if a single byte changes."""
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


# ---------------------------------------------------------------- text extraction


def read_text_file(path: Path) -> str:
    """Read a .txt/.md file whatever encoding Notepad or the user saved it in."""
    raw = path.read_bytes()
    if raw.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")  # UTF-8, with or without a BOM
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")  # older Windows "ANSI" files


def extract_pdf(path: Path) -> tuple[str, list[str]]:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages)
    warnings = []
    if reader.pages and len(text.strip()) < MIN_CHARS_PER_PDF_PAGE * len(reader.pages):
        warnings.append("very little text found; the PDF may be scanned images (text in images is not read)")
    return text, warnings


def extract_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_html(path: Path) -> str:
    soup = BeautifulSoup(read_text_file(path), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    for link in soup.find_all("a", href=True):
        if link["href"].startswith("http"):  # keep where "click here" points to
            link.append(f" ({link['href']})")
    text = soup.get_text(" ")
    return "\n".join(" ".join(line.split()) for line in text.splitlines())


def extract_text(path: Path) -> tuple[str, list[str]]:
    """Return (normalised text, warnings) for any supported document type."""
    ext = path.suffix.lower()
    warnings: list[str] = []
    if ext in {".txt", ".md"}:
        text = read_text_file(path)
    elif ext == ".pdf":
        text, warnings = extract_pdf(path)
    elif ext == ".docx":
        text = extract_docx(path)
    elif ext in {".html", ".htm"}:
        text = extract_html(path)
    else:
        raise IntakeError(f"not a supported document type: {path.name}")
    text = normalise(text)
    if not text.strip():
        warnings.append("no text could be extracted")
    return text, warnings


def normalise(text: str) -> str:
    """Make line endings and spacing consistent so citations (quotes) match reliably later."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [line.rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)  # at most one empty line in a row
    return text.strip() + "\n"


# ---------------------------------------------------------------- installer facts


def read_signer_subject(path: Path) -> str | None:
    """Read which certificate the installer claims to be signed with.

    Uses PowerShell's Get-AuthenticodeSignature, which reads the signature block embedded in the
    file. Only the certificate subject is recorded. Whether the signature is *valid* is a test for
    Agent 2.
    """
    script = "(Get-AuthenticodeSignature -LiteralPath $env:MPAGENT_FILE).SignerCertificate.Subject"
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            env={**os.environ, "MPAGENT_FILE": str(path)},  # pass the path safely, no quoting issues
            capture_output=True,
            text=True,
            timeout=60,
            check=False,  # an unsigned file is not an error; we just get empty output
        )
    except (OSError, subprocess.TimeoutExpired):
        return None
    subject = result.stdout.strip()
    return subject or None


def common_name(subject: str | None) -> str | None:
    """'CN=Kleen Software GmbH, O=..., C=DE' -> 'Kleen Software GmbH'."""
    if not subject:
        return None
    match = re.search(r'CN=("([^"]+)"|[^,]+)', subject)
    if not match:
        return None
    return (match.group(2) or match.group(1)).strip()


def installer_type(path: Path) -> str:
    ext = path.suffix.lower()
    return "zip-bundle" if ext == ".zip" else ext.lstrip(".")


# ---------------------------------------------------------------- main entry


def run_intake(input_dir: Path, runs_root: Path) -> tuple[IntakeManifest, Path]:
    """Run Stage 0 on one plug-in folder. Returns the manifest and the run folder it was saved in."""
    input_dir = input_dir.resolve()
    if not input_dir.is_dir():
        raise IntakeError(f"input folder not found: {input_dir}")

    paths = sorted((p for p in input_dir.iterdir() if p.is_file()), key=lambda p: p.name.lower())
    if not paths:
        raise IntakeError(f"input folder is empty: {input_dir}")

    warnings: list[str] = []
    files: list[InputFile] = []
    for p in paths:
        kind = classify(p)
        files.append(InputFile(name=p.name, kind=kind, size_bytes=p.stat().st_size, sha256=sha256_of(p)))
        if kind is FileKind.UNSUPPORTED:
            warnings.append(f"skipped unsupported file: {p.name}")

    installers = [f for f in files if f.kind is FileKind.INSTALLER]
    if len(installers) != 1:
        found = ", ".join(f.name for f in installers) or "none"
        raise IntakeError(f"expected exactly one installer (.msi/.exe/.zip with a .bundle), found: {found}")
    doc_files = [f for f in files if f.kind is FileKind.DOCUMENT]
    if not doc_files:
        raise IntakeError("no documents found; at least the product description (.txt) is required")

    plugin = input_dir.name
    started = datetime.now().astimezone()  # local time, with the time zone recorded
    run_dir = runs_root.resolve() / plugin / started.strftime("%Y%m%d-%H%M%S")
    (run_dir / "documents").mkdir(parents=True, exist_ok=False)

    documents: list[Document] = []
    total_chars = 0
    for i, f in enumerate(doc_files, start=1):
        src = input_dir / f.name
        text, doc_warnings = extract_text(src)
        doc_id = f"D{i}"
        text_rel = Path("documents") / f"{doc_id}_{src.stem}.txt"
        (run_dir / text_rel).write_text(text, encoding="utf-8")
        total_chars += len(text)
        documents.append(
            Document(
                doc_id=doc_id,
                source_file=f.name,
                format=src.suffix.lower().lstrip("."),
                text_file=text_rel.as_posix(),
                line_count=text.count("\n"),
                char_count=len(text),
                warnings=doc_warnings,
            )
        )
        warnings.extend(f"{doc_id} ({f.name}): {w}" for w in doc_warnings)

    inst = installers[0]
    inst_path = input_dir / inst.name
    subject = read_signer_subject(inst_path)
    if subject is None:
        warnings.append("no signing certificate found in the installer")
    installer = InstallerInfo(
        file=inst.name,
        type=installer_type(inst_path),
        size_bytes=inst.size_bytes,
        sha256=inst.sha256,
        signer_subject=subject,
        signer_name=common_name(subject),
    )

    estimated_tokens = total_chars // 4
    if estimated_tokens > MAX_ESTIMATED_TOKENS:
        warnings.append(
            f"documents are about {estimated_tokens:,} tokens, over the {MAX_ESTIMATED_TOKENS:,} budget; "
            "they may need to be split (not supported yet)"
        )

    metadata = next((f.name for f in files if f.kind is FileKind.METADATA), None)

    manifest = IntakeManifest(
        plugin=plugin,
        input_dir=str(input_dir),
        created_at=started,
        files=files,
        documents=documents,
        installer=installer,
        metadata_file=metadata,
        estimated_tokens=estimated_tokens,
        warnings=warnings,
    )
    (run_dir / "intake_manifest.json").write_text(manifest.model_dump_json(indent=2), encoding="utf-8")
    return manifest, run_dir


def numbered(text: str) -> str:
    """Show text with line numbers, as the LLM will see it: '   3| line text'."""
    return "\n".join(f"{n:>4}| {line}" for n, line in enumerate(text.splitlines(), start=1))
